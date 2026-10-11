"""Tests for the gantt-chart scripts. Run: python3 -m unittest discover -s plugins/pm-toolkit/skills/gantt-chart/tests"""
import datetime as dt
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
SAMPLES = HERE / "sample-plans"
sys.path.insert(0, str(SCRIPTS))

import gantt_core as core  # noqa: E402
from render_html import render_html  # noqa: E402
from render_pptx import render_pptx  # noqa: E402
from render_xlsx import render_xlsx  # noqa: E402

TODAY = dt.date(2026, 11, 20)


def plan_for(name, **kw):
    return core.load_plan(SAMPLES / f"{name}.csv", today=TODAY, **kw)


def by_id(plan):
    return {t["id"]: t for t in plan["tasks"]}


class ScheduleTest(unittest.TestCase):
    def test_small_schedule_and_critical_path(self):
        plan = plan_for("small")
        t = by_id(plan)
        self.assertEqual(plan["unscheduled"], [])
        self.assertEqual((t["T2"]["start"], t["T2"]["end"]), (dt.date(2026, 11, 4), dt.date(2026, 11, 8)))
        self.assertEqual((t["T4"]["start"], t["T4"]["end"]), (dt.date(2026, 11, 16), dt.date(2026, 11, 25)))
        self.assertEqual(t["M1"]["start"], dt.date(2026, 11, 15))
        self.assertEqual(t["M2"]["end"], dt.date(2026, 11, 30))
        self.assertEqual(plan["range"], (dt.date(2026, 11, 2), dt.date(2026, 11, 30)))
        self.assertEqual(plan["critical_path"], ["T1", "T2", "T3", "M1", "T4", "T6", "M2"])
        self.assertFalse(t["T5"]["critical"])
        self.assertEqual(t["T5"]["float_days"], 18)
        self.assertFalse(t["T1"]["derived"])
        self.assertTrue(t["T4"]["derived"])

    def test_missing_dates_are_not_invented(self):
        plan = plan_for("missing-dates")
        t = by_id(plan)
        self.assertEqual([u["id"] for u in plan["unscheduled"]], ["A3", "B3", "C2"])
        for tid in ("A3", "B3", "C2"):
            self.assertIsNone(t[tid]["start"])
            self.assertIsNone(t[tid]["end"])
            self.assertFalse(t[tid]["critical"])
        kinds = {(i["type"], i["id"]) for i in plan["issues"]}
        self.assertIn(("unknown_dependency", "B4"), kinds)
        self.assertIn(("end_before_start", "B2"), kinds)
        self.assertIn(("missing_owner", "A4"), kinds)
        self.assertEqual(t["B2"]["state"], "invalid")
        self.assertEqual(sum(i["type"] == "missing_owner" for i in plan["issues"]), 1)
        self.assertEqual((t["A2"]["start"], t["A2"]["end"]), (dt.date(2026, 11, 14), dt.date(2026, 11, 18)))
        self.assertEqual(t["B4"]["start"], dt.date(2026, 12, 1))

    def test_cycle_flagged_and_left_unscheduled(self):
        tasks, issues = core.normalise([
            {"id": "A", "task": "A", "owner": "x", "start": "2026-01-05", "duration_days": "2"},
            {"id": "B", "task": "B", "owner": "x", "duration_days": "2", "depends_on": "A;C"},
            {"id": "C", "task": "C", "owner": "x", "duration_days": "2", "depends_on": "B"},
        ], "iso")
        plan = core.build_plan(tasks, today=TODAY, issues=issues)
        self.assertTrue(any(i["type"] == "cycle" for i in plan["issues"]))
        self.assertEqual([u["id"] for u in plan["unscheduled"]], ["B", "C"])

    def test_working_days_skip_weekends(self):
        tasks, issues = core.normalise([
            {"id": "A", "task": "A", "owner": "x", "start": "2026-11-05", "duration_days": "3"},
            {"id": "B", "task": "B", "owner": "x", "duration_days": "2", "depends_on": "A"},
        ], "iso")
        plan = core.build_plan(tasks, today=TODAY, day_mode="working", issues=issues)
        t = by_id(plan)
        self.assertEqual(t["A"]["end"], dt.date(2026, 11, 9))
        self.assertEqual((t["B"]["start"], t["B"]["end"]), (dt.date(2026, 11, 10), dt.date(2026, 11, 11)))
        self.assertEqual(plan["critical_path"], ["A", "B"])

    def test_us_style_dates_need_a_decision(self):
        rows = [{"id": "A", "start": "11/16/2026", "end": "11/20/2026"}]
        with self.assertRaises(core.InputError):
            core.detect_date_format(rows)
        self.assertEqual(core.detect_date_format(rows, "mdy")[0], "mdy")
        with self.assertRaises(core.InputError):
            core.detect_date_format([{"start": "02/03/2026"}])

    def test_today_line_only_in_range(self):
        self.assertTrue(plan_for("small")["today_in_range"])
        self.assertFalse(plan_for("large")["today_in_range"])


class RenderTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.out = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_html_is_self_contained(self):
        for name in ("small", "large", "missing-dates"):
            path = render_html(plan_for(name), self.out / f"{name}.html")
            text = path.read_text(encoding="utf-8")
            self.assertIsNone(re.search(r"https?://", text), name)
            self.assertNotRegex(text, r"<(script|link)[^>]+(src|href)=")
            self.assertIn("ctrlKey", text)
        text = (self.out / "missing-dates.html").read_text(encoding="utf-8")
        self.assertIn("Unscheduled, needs dates (3)", text)
        self.assertIn("Persona workshop", text)

    def test_xlsx_conditional_formatting_and_formulas(self):
        from openpyxl import load_workbook
        path = render_xlsx(plan_for("small"), self.out / "small.xlsx")
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Gantt", "Tasks", "Checks"])
        ws = wb["Gantt"]
        self.assertEqual([ws.cell(4, c).value for c in range(1, 10)],
                         ["ID", "Task", "Phase", "Owner", "Start", "End", "Days", "Status", "%"])
        formulas = [f for cf in ws.conditional_formatting for rule in cf.rules for f in rule.formula]
        self.assertGreaterEqual(len(formulas), 5)
        self.assertTrue(any("$E5" in f and "$F5" in f and "L$4" in f for f in formulas))
        self.assertTrue(any('$J5="Yes"' in f for f in formulas), "critical rule")
        self.assertTrue(any('$K5="Yes"' in f for f in formulas), "milestone rule")
        self.assertTrue(any("TODAY()" in f for f in formulas), "today rule")
        days = [ws.cell(r, 7).value for r in range(5, ws.max_row + 1) if isinstance(ws.cell(r, 7).value, str)]
        self.assertTrue(days and all(re.search(r"F\d+-E\d+\+1", d) for d in days))
        self.assertEqual(ws.freeze_panes, "L5")
        self.assertIsInstance(ws.cell(4, 12).value, dt.datetime)
        self.assertEqual(wb["Checks"].cell(2, 1).value, "No issues found")

    def test_xlsx_weekly_for_long_plans(self):
        from openpyxl import load_workbook
        ws = load_workbook(render_xlsx(plan_for("large"), self.out / "large.xlsx"))["Gantt"]
        first, second = ws.cell(4, 12).value, ws.cell(4, 13).value
        self.assertEqual((second - first).days, 7)

    def test_pptx_slide_count(self):
        from pptx import Presentation
        self.assertEqual(len(Presentation(render_pptx(plan_for("small"), self.out / "s.pptx")).slides), 3)
        self.assertEqual(len(Presentation(render_pptx(plan_for("large"), self.out / "l.pptx")).slides), 1 + 3 + 1)

    def test_cli_renders_all_samples(self):
        for name in ("small", "large", "missing-dates"):
            result = subprocess.run(
                [sys.executable, str(SCRIPTS / "gantt.py"), "--input", str(SAMPLES / f"{name}.csv"),
                 "--out-dir", str(self.out), "--name", f"Test {name}", "--today", "2026-11-20"],
                capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            summary = json.loads(result.stdout)
            for ext in ("html", "xlsx", "pptx"):
                self.assertTrue((self.out / f"test-{name}-gantt.{ext}").is_file())
            if name == "missing-dates":
                self.assertEqual(summary["unscheduled_ids"], ["A3", "B3", "C2"])
                self.assertEqual(summary["invalid_ids"], ["B2"])
                self.assertEqual(summary["counts"]["warnings"], 3)
            if name == "small":
                self.assertEqual(summary["critical_path"], ["T1", "T2", "T3", "M1", "T4", "T6", "M2"])

    def test_cli_rejects_bad_input(self):
        result = subprocess.run([sys.executable, str(SCRIPTS / "gantt.py"), "--input", str(self.out / "nope.csv")],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertIn("error", json.loads(result.stdout))


if __name__ == "__main__":
    unittest.main()
