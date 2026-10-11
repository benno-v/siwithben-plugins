#!/usr/bin/env python3
"""Core of the gantt-chart skill: load a task list, schedule it, find the critical path and check it.

Never invents dates or durations. A task is drawn only when its dates are given or can be
derived from its own data (start + duration, end + duration) or from finish-to-start
dependencies plus a duration. Everything else is reported as unscheduled.

Python 3 standard library only (openpyxl is used only to read .xlsx input).
By Benjamin Vermeulen, @SIwithBen. MIT licence.
"""
import csv
import datetime as dt
import json
import re
from pathlib import Path

FIELDS = ["id", "task", "phase", "owner", "start", "end", "duration_days",
          "depends_on", "milestone", "status", "percent_complete"]

ALIASES = {
    "task_id": "id", "key": "id", "wbs": "id",
    "name": "task", "task_name": "task", "title": "task", "summary": "task",
    "workstream": "phase", "group": "phase", "stage": "phase", "epic": "phase",
    "assignee": "owner", "responsible": "owner", "person": "owner",
    "start_date": "start", "begin": "start",
    "end_date": "end", "finish": "end", "finish_date": "end", "due": "end", "due_date": "end",
    "duration": "duration_days", "days": "duration_days",
    "dependencies": "depends_on", "predecessors": "depends_on", "depends": "depends_on",
    "is_milestone": "milestone",
    "percent": "percent_complete", "progress": "percent_complete", "complete": "percent_complete",
    "pct": "percent_complete", "percent_complete_": "percent_complete",
}

DAY_MODES = ("calendar", "working")
DATE_FORMATS = ("auto", "iso", "dmy", "mdy")
YES = {"yes", "y", "true", "1", "x"}
NO = {"no", "n", "false", "0", ""}


class InputError(Exception):
    """The input cannot be read without a decision from the user."""


# ---------- loading ----------

def _norm_header(h):
    h = str(h or "").strip().lower()
    if h in ("%", "% complete", "%complete"):
        return "percent_complete"
    h = re.sub(r"[^a-z0-9]+", "_", h).strip("_")
    return ALIASES.get(h, h)


def _read_csv(path):
    text = Path(path).read_text(encoding="utf-8-sig")
    header = text.splitlines()[0] if text.strip() else ""
    delim = max([",", ";", "\t"], key=header.count) if header else ","
    reader = csv.DictReader(text.splitlines(), delimiter=delim)
    return [{_norm_header(k): v for k, v in row.items() if k is not None} for row in reader]


def _read_xlsx(path):
    try:
        import openpyxl
    except ImportError as exc:
        raise InputError("Reading .xlsx needs openpyxl (pip install openpyxl), or save the sheet as CSV.") from exc
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Tasks"] if "Tasks" in wb.sheetnames else wb.worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    header = [_norm_header(h) for h in rows[0]]
    return [{h: v for h, v in zip(header, r) if h} for r in rows[1:]]


def load_rows(path):
    """Return (rows, meta) from a .json, .csv or .xlsx file. Rows are dicts with normalised keys."""
    path = Path(path)
    if not path.is_file():
        raise InputError(f"Input file not found: {path}")
    suffix = path.suffix.lower()
    meta = {}
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            meta = {k: v for k, v in data.items() if k != "tasks"}
            data = data.get("tasks", [])
        if not isinstance(data, list):
            raise InputError("JSON input must be a list of tasks or an object with a 'tasks' list.")
        rows = [{_norm_header(k): v for k, v in t.items()} for t in data]
    elif suffix in (".csv", ".txt", ".tsv"):
        rows = _read_csv(path)
    elif suffix in (".xlsx", ".xlsm"):
        rows = _read_xlsx(path)
    else:
        raise InputError(f"Unsupported file type '{suffix}'. Use .json, .csv or .xlsx.")
    rows = [r for r in rows if any(v not in (None, "") for v in r.values())]
    return rows, meta


# ---------- dates ----------

SLASH = re.compile(r"^\s*(\d{1,2})[/.](\d{1,2})[/.](\d{4})\s*$")
ISO = re.compile(r"^\s*(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ].*)?$")


def detect_date_format(rows, requested="auto"):
    """Return (format, note). Raises InputError when slash dates clearly are not DD/MM/YYYY."""
    if requested not in DATE_FORMATS:
        raise InputError(f"Unknown date format '{requested}'. Use one of: {', '.join(DATE_FORMATS)}.")
    if requested != "auto":
        return requested, None
    first_big = second_big = seen = False
    for r in rows:
        for key in ("start", "end"):
            m = SLASH.match(str(r.get(key) or ""))
            if m:
                seen = True
                a, b = int(m.group(1)), int(m.group(2))
                first_big |= a > 12
                second_big |= b > 12
    if second_big and first_big:
        raise InputError("Slash dates are mixed: some only work as DD/MM/YYYY and some only as MM/DD/YYYY. "
                         "Fix the source or run again with --date-format dmy or mdy.")
    if second_big:
        raise InputError("Slash dates look like MM/DD/YYYY (e.g. a day above 12 in second place). "
                         "Confirm with the user, then run again with --date-format mdy (or dmy).")
    if not seen:
        return "iso", None
    if not first_big:
        raise InputError("Every slash date could be DD/MM/YYYY or MM/DD/YYYY. I won't guess: "
                         "confirm with the user, then run again with --date-format dmy or mdy.")
    return "dmy", None


def parse_date(value, fmt="dmy"):
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    s = str(value).strip()
    if not s:
        return None
    m = ISO.match(s)
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = SLASH.match(s)
    if m and fmt != "iso":
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        day, month = (b, a) if fmt == "mdy" else (a, b)
        return dt.date(y, month, day)
    raise ValueError(f"unrecognised date '{s}'")


def add_units(d, n, mode):
    """Move n days (calendar) or n working days forward (n >= 0) or backward (n < 0)."""
    if mode == "calendar":
        return d + dt.timedelta(days=n)
    step = 1 if n >= 0 else -1
    d = snap(d, step)
    for _ in range(abs(n)):
        d += dt.timedelta(days=step)
        while d.weekday() >= 5:
            d += dt.timedelta(days=step)
    return d


def snap(d, step=1, mode="working"):
    if mode == "calendar":
        return d
    while d.weekday() >= 5:
        d += dt.timedelta(days=step)
    return d


def unit_index(d, mode):
    """Comparable day number: calendar ordinal or working-day count (weekends share Monday's number)."""
    o = d.toordinal() - 1
    if mode == "calendar":
        return o
    return (o // 7) * 5 + min(o % 7, 5)


def span_units(start, end, mode):
    """Inclusive length of start..end in the chosen unit."""
    if mode == "calendar":
        return (end - start).days + 1
    return unit_index(end + dt.timedelta(days=1), mode) - unit_index(start, mode)


# ---------- normalising ----------

def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def normalise(rows, date_format="dmy"):
    """Turn raw rows into tasks following the task JSON schema. Returns (tasks, issues)."""
    tasks, issues, seen = [], [], {}
    for i, r in enumerate(rows):
        tid = _text(r.get("id"))
        if not tid:
            tid = f"ROW{i + 1}"
            issues.append(_issue("missing_id", tid, f"Row {i + 1} has no ID; labelled {tid}."))
        if tid in seen:
            seen[tid] += 1
            new = f"{tid}-{seen[tid]}"
            issues.append(_issue("duplicate_id", new, f"ID {tid} is used more than once; the repeat is labelled {new}."))
            tid = new
        else:
            seen[tid] = 1
        t = {"id": tid, "task": _text(r.get("task")) or "(no name)", "phase": _text(r.get("phase")),
             "owner": _text(r.get("owner")), "status": _text(r.get("status")), "row": i + 1}
        for key in ("start", "end"):
            try:
                t[key] = parse_date(r.get(key), date_format)
            except ValueError as exc:
                t[key] = None
                issues.append(_issue("bad_date", tid, f"{tid}: {key} date not understood ({exc}); treated as missing."))
        dur = _text(r.get("duration_days"))
        t["duration_days"] = None
        if dur:
            try:
                val = float(dur)
                if val < 0 or not val.is_integer():
                    raise ValueError
                t["duration_days"] = int(val)
            except ValueError:
                issues.append(_issue("bad_duration", tid, f"{tid}: duration '{dur}' is not a whole number of days; treated as missing."))
        deps = r.get("depends_on")
        if isinstance(deps, (list, tuple)):
            deps = [_text(d) for d in deps]
        else:
            deps = re.split(r"[;,]", _text(deps))
        t["depends_on"] = [d.strip() for d in deps if d and d.strip()]
        ms = r.get("milestone")
        t["milestone"] = ms if isinstance(ms, bool) else _text(ms).lower() in YES
        pct = _text(r.get("percent_complete")).rstrip("%").strip()
        t["percent_complete"] = None
        if pct:
            try:
                p = float(pct)
                if 0 < p <= 1 and "." in pct:
                    p *= 100
                t["percent_complete"] = max(0.0, min(100.0, p))
            except ValueError:
                issues.append(_issue("bad_percent", tid, f"{tid}: percent complete '{pct}' is not a number; ignored."))
        tasks.append(t)
    return tasks, issues


def _issue(kind, task_id, message):
    return {"type": kind, "id": task_id, "message": message}


# ---------- scheduling, critical path and checks ----------

def _cycles(ids, edges):
    """Tarjan strongly connected components; returns lists of ids that form cycles."""
    index, low, stack, on, out = {}, {}, [], set(), []
    counter = [0]

    def visit(v):
        index[v] = low[v] = counter[0]
        counter[0] += 1
        stack.append(v)
        on.add(v)
        for w in edges.get(v, []):
            if w not in index:
                visit(w)
                low[v] = min(low[v], low[w])
            elif w in on:
                low[v] = min(low[v], index[w])
        if low[v] == index[v]:
            comp = []
            while True:
                w = stack.pop()
                on.discard(w)
                comp.append(w)
                if w == v:
                    break
            if len(comp) > 1 or v in edges.get(v, []):
                out.append(list(reversed(comp)))

    for v in ids:
        if v not in index:
            visit(v)
    return out


def build_plan(tasks, name="Plan", today=None, day_mode="calendar", issues=None):
    if day_mode not in DAY_MODES:
        raise InputError(f"Unknown day mode '{day_mode}'. Use calendar or working.")
    issues = list(issues or [])
    by_id = {t["id"]: t for t in tasks}
    order = [t["id"] for t in tasks]

    for t in tasks:
        t.update(state="pending", derived=False, reason="", critical=False, float_days=None)
        t["known_deps"] = [d for d in t["depends_on"] if d in by_id and d != t["id"]]
        for d in t["depends_on"]:
            if d not in by_id:
                issues.append(_issue("unknown_dependency", t["id"], f"{t['id']} depends on {d}, which is not in the plan."))
            elif d == t["id"]:
                issues.append(_issue("cycle", t["id"], f"{t['id']} depends on itself."))
        if not t["owner"]:
            issues.append(_issue("missing_owner", t["id"], f"{t['id']} ({t['task']}) has no owner."))

    preds = {t["id"]: t["known_deps"] for t in tasks}
    succ_edges = {i: [] for i in order}
    for t in tasks:
        for d in t["known_deps"]:
            succ_edges[d].append(t["id"])
    cycles = _cycles(order, preds)
    in_cycle = {}
    for n, comp in enumerate(cycles):
        for v in comp:
            in_cycle[v] = n
        issues.append(_issue("cycle", comp[0], "Circular dependency: " + " -> ".join(comp + [comp[0]]) + "."))

    # 1. What each task's own data fixes.
    for t in tasks:
        s, e, dur, ms = t["start"], t["end"], t["duration_days"], t["milestone"]
        if ms and (s or e):
            if s and e and s != e:
                issues.append(_issue("milestone_dates", t["id"], f"{t['id']} is a milestone with different start and end; the end date is used."))
            t["start"] = t["end"] = e or s
            t["state"] = "scheduled"
        elif s and e:
            t["state"] = "scheduled"
            if dur is not None and e >= s and span_units(s, e, day_mode) != dur:
                issues.append(_issue("duration_mismatch", t["id"],
                                     f"{t['id']}: dates span {span_units(s, e, day_mode)} {day_mode} days but duration says {dur}; the dates are used."))
        elif s and dur is not None:
            t["end"] = add_units(snap(s, 1, day_mode), max(dur - 1, 0), day_mode)
            t["state"], t["derived"] = "scheduled", True
        elif e and dur is not None:
            t["start"] = add_units(snap(e, -1, day_mode), -max(dur - 1, 0), day_mode)
            t["state"], t["derived"] = "scheduled", True
        elif t["known_deps"] and (dur is not None or e or ms):
            t["state"] = "pending"
        else:
            t["state"] = "unscheduled"
            if t["depends_on"] and not t["known_deps"]:
                t["reason"] = "Its only dependencies are not in the plan, and it has no dates"
            elif t["known_deps"]:
                t["reason"] = "Has dependencies but no duration or end date"
            elif s:
                t["reason"] = "Has a start date but no end date or duration"
            elif e:
                t["reason"] = "Has an end date but no start date or duration"
            elif dur is not None:
                t["reason"] = "Has a duration but no start date or dependency"
            else:
                t["reason"] = "No start, end, duration or dependency"
        if t["state"] == "scheduled" and t["end"] < t["start"]:
            t["state"] = "invalid"
            issues.append(_issue("end_before_start", t["id"],
                                 f"{t['id']} ends ({t['end'].isoformat()}) before it starts ({t['start'].isoformat()}); not drawn."))

    # 2. Forward pass (finish-to-start) for tasks that need their dependencies.
    changed = True
    while changed:
        changed = False
        for t in tasks:
            if t["state"] != "pending":
                continue
            deps = [by_id[d] for d in t["known_deps"]]
            if any(d["state"] in ("unscheduled", "invalid") for d in deps):
                bad = [d["id"] for d in deps if d["state"] in ("unscheduled", "invalid")]
                t["state"] = "unscheduled"
                t["reason"] = f"Depends on {', '.join(bad)}, which {'is' if len(bad) == 1 else 'are'} not scheduled"
                changed = True
                continue
            if any(d["state"] == "pending" for d in deps):
                continue
            latest = max(d["end"] for d in deps)
            if t["milestone"]:
                t["start"] = t["end"] = latest
            else:
                t["start"] = snap(latest + dt.timedelta(days=1), 1, day_mode)
                if t["duration_days"] is not None:
                    t["end"] = add_units(t["start"], max(t["duration_days"] - 1, 0), day_mode)
            t["state"], t["derived"] = "scheduled", True
            if t["end"] < t["start"]:
                t["state"] = "invalid"
                issues.append(_issue("end_before_start", t["id"],
                                     f"{t['id']}: its dependencies finish on {latest.isoformat()}, after its end date {t['end'].isoformat()}; not drawn."))
            changed = True
    for t in tasks:
        if t["state"] == "pending":
            t["state"] = "unscheduled"
            t["reason"] = "Part of a circular dependency" if t["id"] in in_cycle else "Waiting on dependencies that cannot be scheduled"

    # 3. Explicit dates that clash with dependencies.
    for t in tasks:
        if t["state"] != "scheduled" or t["derived"]:
            continue
        for d in (by_id[x] for x in t["known_deps"]):
            if d["state"] == "scheduled" and (t["start"] < d["end"] or (t["start"] == d["end"] and not t["milestone"])):
                issues.append(_issue("dependency_clash", t["id"],
                                     f"{t['id']} starts {t['start'].isoformat()} but depends on {d['id']}, which ends {d['end'].isoformat()}."))

    # 4. Critical path: zero (or negative) total float, backward pass over the dependency network.
    sched = [t for t in tasks if t["state"] == "scheduled"]
    critical_path = []
    if sched:
        idx = {t["id"]: (unit_index(t["start"], day_mode), unit_index(t["end"], day_mode)) for t in sched}
        project_end = max(ef for _, ef in idx.values())
        ok = set(idx)
        succ = {i: [s for s in succ_edges[i] if s in ok and not (i in in_cycle and in_cycle.get(s) == in_cycle[i])] for i in ok}
        indeg = {i: 0 for i in ok}
        for i in ok:
            for s in succ[i]:
                indeg[s] += 1
        topo, queue = [], [i for i in order if i in ok and indeg[i] == 0]
        while queue:
            v = queue.pop(0)
            topo.append(v)
            for s in succ[v]:
                indeg[s] -= 1
                if indeg[s] == 0:
                    queue.append(s)
        late_start = {}
        for v in reversed(topo):
            es, ef = idx[v]
            lf = project_end
            for s in succ[v]:
                gap = 0 if by_id[s]["milestone"] else 1
                lf = min(lf, late_start[s] - gap)
            late_start[v] = lf - (ef - es)
            by_id[v]["float_days"] = lf - ef
            by_id[v]["critical"] = lf - ef <= 0
        critical_path = [t["id"] for t in sorted(sched, key=lambda t: (t["start"], t["end"], t["row"])) if t["critical"]]

    for t in tasks:
        t.pop("known_deps", None)
    if not any(t["state"] == "scheduled" for t in tasks):
        rng = None
    else:
        rng = (min(t["start"] for t in sched), max(t["end"] for t in sched))
    today = today or dt.date.today()
    unscheduled = [t for t in tasks if t["state"] == "unscheduled"]
    for t in unscheduled:
        issues.append(_issue("unscheduled", t["id"], f"{t['id']} ({t['task']}) is not drawn: {t['reason'].lower()}."))
    return {
        "name": name, "day_mode": day_mode, "today": today,
        "today_in_range": bool(rng and rng[0] <= today <= rng[1]),
        "range": rng, "tasks": tasks, "critical_path": critical_path,
        "unscheduled": unscheduled, "issues": issues, "cycles": cycles,
    }


def load_plan(path, name=None, today=None, day_mode="calendar", date_format="auto"):
    rows, meta = load_rows(path)
    if not rows:
        raise InputError(f"No tasks found in {path}.")
    fmt, note = detect_date_format(rows, date_format)
    tasks, issues = normalise(rows, fmt)
    mode = meta.get("day_mode") if day_mode is None else day_mode
    plan = build_plan(tasks, name or meta.get("name") or Path(path).stem, today, mode or "calendar", issues)
    plan["date_format"] = fmt
    plan["date_note"] = note
    return plan


# ---------- helpers for renderers ----------

def phase_rows(plan, states=("scheduled",)):
    """Yield ('phase', name) and ('task', task) rows, phases in order of first appearance."""
    tasks = [t for t in plan["tasks"] if t["state"] in states]
    phases = []
    for t in tasks:
        if t["phase"] not in phases:
            phases.append(t["phase"])
    use_headers = any(phases)
    rows = []
    for p in phases:
        if use_headers:
            rows.append(("phase", p or "No phase"))
        rows.extend(("task", t) for t in tasks if t["phase"] == p)
    return rows


ISSUE_LABELS = {
    "unscheduled": "Unscheduled, needs dates", "cycle": "Circular dependency",
    "unknown_dependency": "Unknown dependency", "end_before_start": "End before start",
    "missing_owner": "Missing owner", "duplicate_id": "Duplicate ID", "missing_id": "Missing ID",
    "bad_date": "Date not understood", "bad_duration": "Duration not understood",
    "bad_percent": "Percent not understood", "duration_mismatch": "Duration does not match dates",
    "dependency_clash": "Starts before a dependency finishes", "milestone_dates": "Milestone dates",
}


def task_json(t):
    out = {k: t.get(k) for k in FIELDS}
    for k in ("start", "end"):
        out[k] = t[k].isoformat() if t.get(k) else None
    out.update(state=t["state"], derived=t["derived"], critical=t["critical"],
               float_days=t["float_days"], reason=t["reason"] or None)
    return out


def summary(plan, files=()):
    tasks = plan["tasks"]
    rng = plan["range"]
    return {
        "name": plan["name"],
        "files": [str(f) for f in files],
        "day_mode": plan["day_mode"],
        "date_format": plan.get("date_format"),
        "date_note": plan.get("date_note"),
        "range": [rng[0].isoformat(), rng[1].isoformat()] if rng else None,
        "counts": {
            "tasks": len(tasks),
            "scheduled": sum(t["state"] == "scheduled" for t in tasks),
            "derived": sum(t["state"] == "scheduled" and t["derived"] for t in tasks),
            "unscheduled": len(plan["unscheduled"]),
            "invalid": sum(t["state"] == "invalid" for t in tasks),
            "milestones": sum(t["milestone"] for t in tasks),
            "critical": len(plan["critical_path"]),
            "warnings": sum(i["type"] != "unscheduled" for i in plan["issues"]),
        },
        "critical_path": plan["critical_path"],
        "unscheduled_ids": [t["id"] for t in plan["unscheduled"]],
        "invalid_ids": [t["id"] for t in tasks if t["state"] == "invalid"],
        "warnings": [i["message"] for i in plan["issues"] if i["type"] != "unscheduled"],
    }
