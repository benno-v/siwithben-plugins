#!/usr/bin/env python3
"""Render a plan as an Excel workbook whose Gantt bars are conditional formatting rules,
so they move when the user edits Start or End. Needs openpyxl.

By Benjamin Vermeulen, @SIwithBen. MIT licence.
"""
import datetime as dt

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from gantt_core import ISSUE_LABELS, phase_rows

NAVY, TEAL, AMBER, RED = "1B2A4A", "13838F", "F0A202", "C8423B"
BAND, LINE, MUTED = "F3F6FA", "DFE4EC", "5B6578"
HEAD_ROW, FIRST_ROW = 4, 5
COLS = ["ID", "Task", "Phase", "Owner", "Start", "End", "Days", "Status", "%", "Critical", "Milestone"]
WIDTHS = [8, 34, 16, 18, 12, 12, 7, 13, 6, 9, 10]
FIRST_DATE_COL = len(COLS) + 1
DATE_FMT = "dd mmm yyyy"


def _fill(hex_colour):
    return PatternFill("solid", start_color=hex_colour, end_color=hex_colour)


def _write_task_row(ws, r, t, mode):
    ws.cell(r, 1, t["id"])
    ws.cell(r, 2, t["task"])
    ws.cell(r, 3, t["phase"])
    ws.cell(r, 4, t["owner"] or None)
    for c, key in ((5, "start"), (6, "end")):
        cell = ws.cell(r, c, t[key])
        cell.number_format = DATE_FMT
    days = f"NETWORKDAYS(E{r},F{r})" if mode == "working" else f"F{r}-E{r}+1"
    ws.cell(r, 7, f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),{days},"")').alignment = Alignment(horizontal="center")
    ws.cell(r, 8, t["status"] or None)
    pct = ws.cell(r, 9, t["percent_complete"] / 100 if t["percent_complete"] is not None else None)
    pct.number_format = "0%"
    ws.cell(r, 10, "Yes" if t["critical"] else "No")
    ws.cell(r, 11, "Yes" if t["milestone"] else "No")
    if not t["owner"]:
        ws.cell(r, 4).fill = _fill("FFF4DC")


def render_xlsx(plan, out_path):
    mode = plan["day_mode"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt"
    rng = plan["range"]
    span = (rng[1] - rng[0]).days + 1 if rng else 0
    weekly = span > 120
    if rng:
        first = rng[0] - dt.timedelta(days=rng[0].weekday()) if weekly else rng[0]
        step = 7 if weekly else 1
        dates = []
        d = first
        while d <= rng[1]:
            dates.append(d)
            d += dt.timedelta(days=step)
    else:
        dates = []
    last_date_col = FIRST_DATE_COL + len(dates) - 1

    unit = "working days" if mode == "working" else "calendar days"
    ws["A1"] = f"{plan['name']}: Gantt chart"
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws["A2"] = (f"Durations in {unit}. {'Weekly' if weekly else 'Daily'} columns. Teal = task, amber = critical path, "
                f"\u25c6 = milestone, red column = today. Edit Start or End and the bars move. Critical flags were "
                f"calculated when the file was made; set Critical or Milestone to Yes or No to change the colours.")
    ws["A2"].font = Font(size=9, italic=True, color=MUTED)

    thin = Side(style="thin", color=LINE)
    for c, (label, width) in enumerate(zip(COLS, WIDTHS), start=1):
        cell = ws.cell(HEAD_ROW, c, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _fill(NAVY)
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    last_month = None
    for i, d in enumerate(dates):
        c = FIRST_DATE_COL + i
        cell = ws.cell(HEAD_ROW, c, d)
        cell.number_format = "dd" if not weekly else "dd mmm"
        cell.font = Font(bold=True, color="FFFFFF", size=8)
        cell.fill = _fill(NAVY)
        cell.alignment = Alignment(horizontal="center", text_rotation=0 if not weekly else 90)
        ws.column_dimensions[get_column_letter(c)].width = 3.2 if not weekly else 3.6
        if (d.year, d.month) != last_month:
            m = ws.cell(HEAD_ROW - 1, c, d.strftime("%b %Y"))
            m.font = Font(bold=True, color=NAVY, size=9)
            last_month = (d.year, d.month)
    if weekly:
        ws.row_dimensions[HEAD_ROW].height = 40

    r = FIRST_ROW
    task_rows = []
    for kind, item in phase_rows(plan):
        if kind == "phase":
            ws.cell(r, 1, item).font = Font(bold=True, color=NAVY)
            for c in range(1, max(last_date_col, len(COLS)) + 1):
                ws.cell(r, c).fill = _fill(BAND)
        else:
            _write_task_row(ws, r, item, mode)
            task_rows.append(r)
        r += 1
    others = [t for t in plan["tasks"] if t["state"] != "scheduled"]
    if others:
        ws.cell(r, 1, "Unscheduled, needs dates (add Start and End and the bars appear)").font = Font(bold=True, color="7A5200")
        for c in range(1, max(last_date_col, len(COLS)) + 1):
            ws.cell(r, c).fill = _fill("FFF4DC")
        r += 1
        for t in others:
            _write_task_row(ws, r, t, mode)
            task_rows.append(r)
            r += 1
    last_row = r - 1

    for tr in task_rows:
        for c in range(1, len(COLS) + 1):
            ws.cell(tr, c).border = Border(bottom=thin)

    if dates and task_rows:
        fc = get_column_letter(FIRST_DATE_COL)
        lc = get_column_letter(last_date_col)
        f, h = FIRST_ROW, HEAD_ROW
        within = (f"{fc}${h}>=$E{f},{fc}${h}<=$F{f}" if not weekly
                  else f"{fc}${h}<=$F{f},{fc}${h}+6>=$E{f}")
        ms_on = (f"$F{f}={fc}${h}" if not weekly else f"$F{f}>={fc}${h},$F{f}<={fc}${h}+6")
        dated = f"ISNUMBER($E{f}),ISNUMBER($F{f})"
        for tr in task_rows:
            for c in range(FIRST_DATE_COL, last_date_col + 1):
                col = get_column_letter(c)
                on = (f"$F{tr}={col}${h}" if not weekly else f"$F{tr}>={col}${h},$F{tr}<={col}${h}+6")
                cell = ws.cell(tr, c, f'=IF(AND($K{tr}="Yes",ISNUMBER($F{tr}),{on}),"\u25c6","")')
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(size=9)
        area = f"{fc}{FIRST_ROW}:{lc}{last_row}"
        today = (f"{fc}${h}=TODAY()" if not weekly else f"AND(TODAY()>={fc}${h},TODAY()<={fc}${h}+6)")
        rules = [
            (area, FormulaRule(formula=[f'AND($K{f}="Yes",$J{f}="Yes",{dated},{ms_on})'],
                               font=Font(bold=True, color="B97A00"), stopIfTrue=True)),
            (area, FormulaRule(formula=[f'AND($K{f}="Yes",{dated},{ms_on})'],
                               font=Font(bold=True, color=NAVY), stopIfTrue=True)),
            (area, FormulaRule(formula=[f'AND($K{f}<>"Yes",$J{f}="Yes",{dated},{within})'],
                               fill=_fill(AMBER), font=Font(color=AMBER))),
            (area, FormulaRule(formula=[f'AND($K{f}<>"Yes",{dated},{within})'],
                               fill=_fill(TEAL), font=Font(color=TEAL))),
            (f"{fc}{HEAD_ROW}:{lc}{last_row}", FormulaRule(formula=[today], fill=_fill("F6D5D2"),
                                                           border=Border(left=Side(style="thin", color=RED), right=Side(style="thin", color=RED)))),
        ]
        for ref, rule in rules:
            ws.conditional_formatting.add(ref, rule)

    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(yes_no)
    if task_rows:
        yes_no.add(f"J{FIRST_ROW}:K{last_row}")
    ws.freeze_panes = ws.cell(FIRST_ROW, FIRST_DATE_COL)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEAD_ROW - 1}:{HEAD_ROW}"

    # Tasks: the normalised data, one row per task in input order.
    ts = wb.create_sheet("Tasks")
    head = ["ID", "Task", "Phase", "Owner", "Start", "End", "Duration (days)", "Depends on", "Milestone",
            "Status", "% complete", "Drawn", "Dates derived", "Critical", f"Total float ({unit})", "Note"]
    ts.append(head)
    for c in range(1, len(head) + 1):
        ts.cell(1, c).font = Font(bold=True, color="FFFFFF")
        ts.cell(1, c).fill = _fill(NAVY)
    for t in plan["tasks"]:
        note = t["reason"] if t["state"] == "unscheduled" else ("End before start" if t["state"] == "invalid" else "")
        ts.append([t["id"], t["task"], t["phase"], t["owner"], t["start"], t["end"], t["duration_days"],
                   "; ".join(t["depends_on"]), "Yes" if t["milestone"] else "No", t["status"], t["percent_complete"],
                   "Yes" if t["state"] == "scheduled" else "No", "Yes" if t["derived"] else "No",
                   "Yes" if t["critical"] else "No", t["float_days"], note])
        for c in (5, 6):
            ts.cell(ts.max_row, c).number_format = DATE_FMT
    for c, w in enumerate([8, 34, 16, 18, 13, 13, 10, 14, 10, 13, 10, 8, 10, 9, 12, 44], start=1):
        ts.column_dimensions[get_column_letter(c)].width = w
    ts.freeze_panes = "C2"
    ts.auto_filter.ref = f"A1:{get_column_letter(len(head))}{ts.max_row}"

    # Checks: every validation issue, including unscheduled tasks.
    cs = wb.create_sheet("Checks")
    cs.append(["Check", "Task ID", "Detail"])
    for c in range(1, 4):
        cs.cell(1, c).font = Font(bold=True, color="FFFFFF")
        cs.cell(1, c).fill = _fill(NAVY)
    for i in plan["issues"]:
        cs.append([ISSUE_LABELS.get(i["type"], i["type"]), i["id"], i["message"]])
    if plan.get("date_note"):
        cs.append(["Date format", "", plan["date_note"]])
    if cs.max_row == 1:
        cs.append(["No issues found", "", "Every task has dates, an owner and known dependencies."])
    cs.column_dimensions["A"].width = 32
    cs.column_dimensions["B"].width = 10
    cs.column_dimensions["C"].width = 100

    wb.save(out_path)
    return out_path
