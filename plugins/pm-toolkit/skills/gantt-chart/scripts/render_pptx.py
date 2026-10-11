#!/usr/bin/env python3
"""Render a plan as a 16:9 PowerPoint deck: title, Gantt slides drawn with shapes, gaps and checks.
Needs python-pptx.

By Benjamin Vermeulen, @SIwithBen. MIT licence.
"""
import datetime as dt

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from gantt_core import ISSUE_LABELS, phase_rows

NAVY, TEAL, AMBER, RED = RGBColor(0x1B, 0x2A, 0x4A), RGBColor(0x13, 0x83, 0x8F), RGBColor(0xF0, 0xA2, 0x02), RGBColor(0xC8, 0x42, 0x3B)
AMBER_D, BAND, LINE, MUTED, WHITE = RGBColor(0xB9, 0x7A, 0x00), RGBColor(0xF3, 0xF6, 0xFA), RGBColor(0xDF, 0xE4, 0xEC), RGBColor(0x5B, 0x65, 0x78), RGBColor(0xFF, 0xFF, 0xFF)
TASKS_PER_SLIDE = 18
CHECK_LINES_PER_SLIDE = 16
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
LEFT, LABEL_W, TOP_AXIS, BOTTOM = Inches(0.4), Inches(4.1), Inches(1.25), Inches(6.85)


def _fmt(d):
    return f"{d.day} {d.strftime('%b %Y')}" if d else "Not set"


def _box(slide, x, y, w, h, colour, shape=MSO_SHAPE.RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(max(w, 1))), Emu(int(max(h, 1))))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = colour
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    return s


def _text(slide, x, y, w, h, text, size=11, bold=False, colour=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for n, line in enumerate(lines):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = colour
    return tb


def _line(slide, x1, y1, x2, y2, colour, width=1.0, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x1)), Emu(int(y1)), Emu(int(x2)), Emu(int(y2)))
    c.line.color.rgb = colour
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return c


def _title_bar(slide, title, subtitle=""):
    _box(slide, 0, 0, SLIDE_W, Inches(0.9), NAVY)
    _text(slide, LEFT, Inches(0.12), Inches(10), Inches(0.45), title, size=22, bold=True, colour=WHITE)
    if subtitle:
        _text(slide, LEFT, Inches(0.52), Inches(12), Inches(0.3), subtitle, size=11, colour=RGBColor(0xC9, 0xD3, 0xE6))


def _pages(rows):
    """Split phase/task rows into pages of at most TASKS_PER_SLIDE tasks, repeating the phase header."""
    pages, cur, n, phase = [], [], 0, None
    for kind, item in rows:
        if kind == "phase":
            phase = item
            if n >= TASKS_PER_SLIDE:
                pages.append(cur)
                cur, n = [], 0
            cur.append((kind, item))
            continue
        if n >= TASKS_PER_SLIDE:
            pages.append(cur)
            cur, n = ([("phase", phase)] if phase is not None else []), 0
        cur.append((kind, item))
        n += 1
    if cur:
        pages.append(cur)
    return pages


def _gantt_slide(prs, plan, rows, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rng = plan["range"]
    unit = "working days" if plan["day_mode"] == "working" else "calendar days"
    _title_bar(slide, f"{plan['name']}: Gantt chart" + (f" ({page} of {total})" if total > 1 else ""),
               f"{_fmt(rng[0])} to {_fmt(rng[1])}. Durations in {unit}.")
    chart_x = LEFT + LABEL_W
    chart_w = SLIDE_W - chart_x - Inches(0.4)
    span = (rng[1] - rng[0]).days + 1
    d0, d1 = rng[0], rng[1] + dt.timedelta(days=1 + max(1, round(span * 0.04)))
    days = (d1 - d0).days
    per_day = chart_w / days

    def xpos(d):
        return chart_x + (d - d0).days * per_day

    axis_y = TOP_AXIS
    body_top = axis_y + Inches(0.35)
    row_h = min(Inches(0.3), (BOTTOM - body_top) / max(len(rows), 1))

    # Axis: weekly ticks for short plans, monthly for longer ones.
    _line(slide, chart_x, body_top, chart_x + chart_w, body_top, MUTED, 0.75)
    ticks = []
    if days <= 70:
        d = d0 + dt.timedelta(days=(7 - d0.weekday()) % 7)
        while d < d1:
            ticks.append((d, f"{d.day} {d.strftime('%b')}"))
            d += dt.timedelta(days=7)
    else:
        d = dt.date(d0.year, d0.month, 1)
        while d < d1:
            if d >= d0:
                ticks.append((d, d.strftime("%b %Y") if d.month == 1 or not ticks else d.strftime("%b")))
            d = dt.date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    for d, label in ticks:
        x = xpos(d)
        _line(slide, x, body_top, x, body_top + row_h * len(rows), LINE, 0.5)
        _text(slide, x, axis_y, Inches(0.9), Inches(0.3), label, size=9, colour=MUTED)
    _text(slide, LEFT, axis_y, Inches(0.6), Inches(0.3), "ID", size=9, bold=True, colour=MUTED)
    _text(slide, LEFT + Inches(0.6), axis_y, Inches(2.3), Inches(0.3), "Task", size=9, bold=True, colour=MUTED)
    _text(slide, LEFT + Inches(2.9), axis_y, Inches(1.2), Inches(0.3), "Owner", size=9, bold=True, colour=MUTED)

    size = 9 if row_h >= Inches(0.24) else 8
    for i, (kind, item) in enumerate(rows):
        y = body_top + i * row_h
        if kind == "phase":
            _box(slide, LEFT, y, SLIDE_W - LEFT - Inches(0.4), row_h, BAND)
            _text(slide, LEFT, y, LABEL_W, row_h, item, size=size, bold=True)
            continue
        t = item
        _text(slide, LEFT, y, Inches(0.6), row_h, t["id"], size=size, colour=MUTED)
        name = t["task"] if len(t["task"]) <= 38 else t["task"][:36] + "..."
        _text(slide, LEFT + Inches(0.6), y, Inches(2.3), row_h, name, size=size, bold=t["critical"], colour=NAVY)
        owner = t["owner"] or "No owner"
        _text(slide, LEFT + Inches(2.9), y, Inches(1.2), row_h, owner if len(owner) <= 18 else owner[:16] + "...",
              size=size - 1, colour=MUTED if t["owner"] else AMBER_D)
        colour = AMBER if t["critical"] else TEAL
        if t["milestone"]:
            s = row_h * 0.62
            cx = xpos(t["end"] + dt.timedelta(days=1))
            _box(slide, cx - s / 2, y + (row_h - s) / 2, s, s, AMBER_D if t["critical"] else NAVY, MSO_SHAPE.DIAMOND)
        else:
            x = xpos(t["start"])
            w = xpos(t["end"] + dt.timedelta(days=1)) - x
            _box(slide, x, y + row_h * 0.2, w, row_h * 0.6, colour, MSO_SHAPE.RECTANGLE)
            if t["percent_complete"]:
                _box(slide, x, y + row_h * 0.68, w * min(t["percent_complete"], 100) / 100, row_h * 0.12, NAVY)
    if plan["today_in_range"]:
        tx = xpos(plan["today"]) + per_day / 2
        _line(slide, tx, axis_y + Inches(0.25), tx, body_top + row_h * len(rows), RED, 1.5, dash=True)
        _text(slide, tx + Inches(0.03), body_top - Inches(0.02), Inches(0.6), Inches(0.2), "Today", size=8, bold=True, colour=RED)

    # Legend.
    ly, lx = Inches(7.0), LEFT
    for label, colour, shape in (("Task", TEAL, MSO_SHAPE.RECTANGLE), ("Critical path", AMBER, MSO_SHAPE.RECTANGLE),
                                 ("Milestone", NAVY, MSO_SHAPE.DIAMOND), ("Today", RED, None)):
        if shape is None:
            _line(slide, lx, ly + Inches(0.1), lx + Inches(0.3), ly + Inches(0.1), colour, 1.5, dash=True)
        else:
            w = Inches(0.3) if shape == MSO_SHAPE.RECTANGLE else Inches(0.18)
            _box(slide, lx, ly + Inches(0.02), w, Inches(0.16), colour, shape)
        _text(slide, lx + Inches(0.36), ly, Inches(1.3), Inches(0.2), label, size=9, colour=MUTED)
        lx += Inches(1.6)
    n_uns = len(plan["unscheduled"])
    if n_uns:
        _text(slide, lx + Inches(0.4), ly, Inches(5), Inches(0.2),
              f"{n_uns} unscheduled task{'s' if n_uns != 1 else ''} not drawn: see Gaps and checks.", size=9, colour=AMBER_D)
    return slide


def _check_lines(plan):
    lines = []
    if plan["unscheduled"]:
        lines.append(("head", f"Unscheduled, needs dates ({len(plan['unscheduled'])})"))
        lines += [("item", f"{t['id']}  {t['task']}: {t['reason']}.") for t in plan["unscheduled"]]
    warns = [i for i in plan["issues"] if i["type"] != "unscheduled"]
    if warns:
        lines.append(("head", f"Checks ({len(warns)})"))
        lines += [("item", f"{ISSUE_LABELS.get(i['type'], i['type'])}: {i['message']}") for i in warns]
    if plan.get("date_note"):
        lines.append(("item", plan["date_note"]))
    return lines


def render_pptx(plan, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    tasks = plan["tasks"]
    rng = plan["range"]
    unit = "working days" if plan["day_mode"] == "working" else "calendar days"

    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    _box(s, Inches(0.8), Inches(3.55), Inches(1.2), Inches(0.08), TEAL)
    _text(s, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.2), plan["name"], size=40, bold=True, colour=WHITE, anchor=MSO_ANCHOR.BOTTOM)
    _text(s, Inches(0.8), Inches(3.8), Inches(11.5), Inches(0.5), "Gantt chart", size=22, colour=RGBColor(0xC9, 0xD3, 0xE6))
    drawn = sum(t["state"] == "scheduled" for t in tasks)
    facts = [f"{_fmt(rng[0])} to {_fmt(rng[1])}" if rng else "No scheduled tasks yet",
             f"{drawn} of {len(tasks)} tasks drawn, {sum(t['milestone'] and t['state'] == 'scheduled' for t in tasks)} milestones, "
             f"{len(plan['critical_path'])} on the critical path, {len(plan['unscheduled'])} unscheduled",
             f"Durations in {unit}. Generated {_fmt(plan['today'])}. Draft for review."]
    _text(s, Inches(0.8), Inches(4.5), Inches(11.5), Inches(1.2), facts, size=14, colour=WHITE, anchor=MSO_ANCHOR.TOP)

    if rng:
        pages = _pages(phase_rows(plan))
        for n, rows in enumerate(pages, start=1):
            _gantt_slide(prs, plan, rows, n, len(pages))
    else:
        g = prs.slides.add_slide(prs.slide_layouts[6])
        _title_bar(g, f"{plan['name']}: Gantt chart")
        _text(g, LEFT, Inches(1.5), Inches(12), Inches(0.5), "No scheduled tasks to draw yet. See Gaps and checks.", size=16, colour=MUTED)

    lines = _check_lines(plan) or [("item", "No gaps or issues found: every task has dates, an owner and known dependencies.")]
    chunks = [lines[i:i + CHECK_LINES_PER_SLIDE] for i in range(0, len(lines), CHECK_LINES_PER_SLIDE)]
    for n, chunk in enumerate(chunks, start=1):
        c = prs.slides.add_slide(prs.slide_layouts[6])
        _title_bar(c, "Gaps and checks" + (f" ({n} of {len(chunks)})" if len(chunks) > 1 else ""),
                   "Fix these in the source plan, then run the chart again. Nothing here was guessed.")
        tb = c.shapes.add_textbox(LEFT, Inches(1.2), SLIDE_W - 2 * LEFT, Inches(5.9))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, (kind, text) in enumerate(chunk):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = text if kind == "head" else f"\u2022  {text}"
            run.font.size = Pt(15 if kind == "head" else 12)
            run.font.bold = kind == "head"
            run.font.color.rgb = NAVY if kind == "head" else RGBColor(0x1D, 0x24, 0x33)
            p.space_before = Pt(10 if kind == "head" and i else 2)
    prs.save(out_path)
    return out_path
